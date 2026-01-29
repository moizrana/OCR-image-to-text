"""
Image to Text OCR Application
Powered by Google Gemini API
"""
import streamlit as st
from PIL import Image
import io
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

from gemini_ocr import GeminiOCR


# Page configuration
st.set_page_config(
    page_title="Image to Text OCR",
    page_icon="🔍",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Custom CSS for modern styling
st.markdown("""
    <style>
    .main {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        padding: 2rem;
    }
    .stApp {
        background: transparent;
    }
    .upload-container {
        background: white;
        padding: 2rem;
        border-radius: 1rem;
        box-shadow: 0 8px 32px rgba(0, 0, 0, 0.1);
    }
    .result-container {
        background: white;
        padding: 2rem;
        border-radius: 1rem;
        box-shadow: 0 8px 32px rgba(0, 0, 0, 0.1);
        margin-top: 2rem;
    }
    h1 {
        color: white;
        text-align: center;
        font-size: 3rem;
        font-weight: bold;
        margin-bottom: 1rem;
        text-shadow: 2px 2px 4px rgba(0, 0, 0, 0.2);
    }
    .subtitle {
        color: white;
        text-align: center;
        font-size: 1.2rem;
        margin-bottom: 2rem;
        opacity: 0.9;
    }
    .stButton > button {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        color: white;
        border: none;
        padding: 0.75rem 2rem;
        border-radius: 0.5rem;
        font-weight: bold;
        transition: transform 0.2s;
    }
    .stButton > button:hover {
        transform: scale(1.05);
    }
    </style>
""", unsafe_allow_html=True)

# Title
st.markdown("<h1>🔍 Image to Text OCR</h1>", unsafe_allow_html=True)
st.markdown("<p class='subtitle'>Extract text from images using Google Gemini AI</p>", unsafe_allow_html=True)

# Sidebar
with st.sidebar:
    st.header("⚙️ Configuration")
    
    # API Key input
    api_key = os.getenv("GEMINI_API_KEY", "")
    
    if not api_key:
        api_key = st.text_input(
            "Gemini API Key",
            type="password",
            help="Enter your Google Gemini API key. Get one at https://makersuite.google.com/app/apikey",
            value="AIzaSyBaW0Lo4ROcleus_T1qJNKZp6zdFzAUccM"
        )
    
    # Model selection
    model_choice = st.selectbox(
        "Model",
        ["gemini-2.5-flash", "gemini-2.5-pro"],
        index=0,  # Default to gemini-2.5-flash (faster)
        help="Flash is faster and great for most tasks. Pro is more accurate for complex documents."
    )
    
    # Extraction mode
    extraction_mode = st.radio(
        "Extraction Mode",
        ["Simple", "Structured"],
        help="Simple: Extract text only. Structured: Include document analysis"
    )
    
    st.divider()
    
    st.info("📝 **Tip**: This app works best with clear, well-lit images of documents, handwritten notes, or printed text.")
    
    st.divider()
    
    # About section
    with st.expander("ℹ️ About"):
        st.markdown("""
        **Image to Text OCR**
        
        Powered by Google Gemini Vision API
        
        Features:
        - 🤖 AI-powered text extraction
        - 📄 Support for printed and handwritten text
        - 💾 Export to Word documents
        - 🎨 Preserves formatting and structure
        
        Version 2.0
        """)

# Main content
if not api_key:
    st.warning("⚠️ Please enter your Gemini API key in the sidebar to get started.")
    st.stop()

# Initialize session state
if 'extracted_text' not in st.session_state:
    st.session_state.extracted_text = None
if 'structure_info' not in st.session_state:
    st.session_state.structure_info = None
if 'uploaded_image' not in st.session_state:
    st.session_state.uploaded_image = None

# File upload
col1, col2 = st.columns([1, 1])

with col1:
    st.markdown("<div class='upload-container'>", unsafe_allow_html=True)
    st.subheader("📤 Upload Image")
    
    uploaded_file = st.file_uploader(
        "Choose an image file",
        type=['png', 'jpg', 'jpeg', 'webp'],
        help="Supported formats: PNG, JPG, JPEG, WEBP"
    )
    
    if uploaded_file:
        # Display uploaded image
        image = Image.open(uploaded_file)
        st.image(image, caption="Uploaded Image", use_container_width=True)
        st.session_state.uploaded_image = image
        
        # Extract button
        if st.button("🚀 Extract Text", use_container_width=True):
            with st.spinner("🔄 Processing image with Gemini AI..."):
                try:
                    # Initialize OCR service
                    ocr = GeminiOCR(api_key=api_key, model_name=model_choice)
                    
                    # Extract text
                    if extraction_mode == "Simple":
                        text = ocr.extract_text(image)
                        st.session_state.extracted_text = text
                        st.session_state.structure_info = None
                    else:
                        result = ocr.extract_text_structured(image)
                        st.session_state.extracted_text = result['text']
                        st.session_state.structure_info = result['structure']
                    
                    st.success("✅ Text extracted successfully!")
                    st.rerun()
                    
                except Exception as e:
                    st.error(f"❌ Error: {str(e)}")
    
    st.markdown("</div>", unsafe_allow_html=True)

with col2:
    st.markdown("<div class='upload-container'>", unsafe_allow_html=True)
    st.subheader("📝 Extracted Text")
    
    if st.session_state.extracted_text:
        # Display structure info if available
        if st.session_state.structure_info:
            with st.expander("📊 Document Analysis", expanded=False):
                st.text(st.session_state.structure_info)
        
        # Display extracted text
        text_area = st.text_area(
            "Extracted Text",
            st.session_state.extracted_text,
            height=400,
            help="Edit the text if needed before downloading"
        )
        
        # Update extracted text if edited
        st.session_state.extracted_text = text_area
        
        # Statistics
        col_a, col_b, col_c = st.columns(3)
        with col_a:
            st.metric("Characters", len(text_area))
        with col_b:
            st.metric("Words", len(text_area.split()))
        with col_c:
            st.metric("Lines", len(text_area.split('\n')))
        
        st.divider()
        
        # Download buttons
        col_x, col_y = st.columns(2)
        
        with col_x:
            # Download as text
            st.download_button(
                label="📄 Download as TXT",
                data=text_area,
                file_name=f"extracted_text_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt",
                mime="text/plain",
                use_container_width=True
            )
        
        with col_y:
            # Download as Word
            if st.button("📘 Download as DOCX", use_container_width=True):
                # Create Word document
                doc = Document()
                
                # Add title
                title = doc.add_heading('Extracted Text from OCR', 0)
                title.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
                # Add metadata
                doc.add_paragraph(f"Extracted on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
                doc.add_paragraph(f"Model used: {model_choice}")
                doc.add_paragraph("")
                
                # Add extracted text
                for paragraph in text_area.split('\n\n'):
                    if paragraph.strip():
                        p = doc.add_paragraph(paragraph.strip())
                        p.style.font.size = Pt(11)
                
                # Save to buffer
                buffer = io.BytesIO()
                doc.save(buffer)
                buffer.seek(0)
                
                # Download button
                st.download_button(
                    label="⬇️ Click to Download DOCX",
                    data=buffer,
                    file_name=f"extracted_text_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True
                )
    else:
        st.info("👆 Upload an image and click 'Extract Text' to see results here.")
    
    st.markdown("</div>", unsafe_allow_html=True)

# Footer
st.markdown("---")
st.markdown(
    "<p style='text-align: center; color: white; opacity: 0.7;'>Built with ❤️ using Streamlit and Google Gemini AI</p>",
    unsafe_allow_html=True
)
